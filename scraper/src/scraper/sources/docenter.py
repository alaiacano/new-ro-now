"""
DocumentCenter mirror: discover all folders/documents, then download them locally.

Discovery uses Playwright because:
  - The folder tree is a JS-rendered Ant Design component
  - The document list is loaded via an authenticated AJAX API (POST)
  - Direct httpx calls to the API return 404 without a valid browser session

Two-phase process:
  scrape discover   -- Playwright clicks each folder → captures folderId from POST body
                       → paginates API via page.evaluate(fetch()) → populates SQLite
  scrape download   -- httpx downloads each file → data/doc_cache/{id}.ext

Idempotent: discover uses INSERT OR IGNORE; download skips already-downloaded files.
"""
import asyncio
import json
import mimetypes
import re
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlparse

import httpx
import pypdf
from playwright.async_api import async_playwright
from rich.console import Console
from rich.progress import BarColumn, MofNCompleteColumn, Progress, SpinnerColumn, TextColumn

from ..config import BASE_URL, DOCENTER_URL, DOC_CACHE_DIR
from ..db import get_conn, init_db, log_step

console = Console()

HEADERS = {"User-Agent": "new-ro-monitor/1.0 (personal civic research tool)"}
API_PATH = "/Admin/DocumentCenter/Home/Document_AjaxBinding?renderMode=0&loadSource=7"
ROWS_PER_PAGE = 100
DOWNLOAD_CONCURRENCY = 3
DOWNLOAD_DELAY_S = 0.5


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _now() -> str:
    return datetime.now(tz=timezone.utc).isoformat()


def _ext_from_content_type(ct: str) -> str:
    ct = ct.split(";")[0].strip().lower()
    mapping = {
        "application/pdf": ".pdf",
        "application/msword": ".doc",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
        "application/vnd.ms-excel": ".xls",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx",
        "application/vnd.ms-powerpoint": ".ppt",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation": ".pptx",
        "image/png": ".png",
        "image/jpeg": ".jpg",
        "image/gif": ".gif",
        "image/svg+xml": ".svg",
        "text/plain": ".txt",
    }
    return mapping.get(ct, mimetypes.guess_extension(ct) or ".bin")


def _ext_from_url(url: str) -> str:
    suffix = Path(urlparse(url).path).suffix.lower()
    return suffix if suffix else ""


def _extract_pdf_text(path: Path) -> str:
    reader = pypdf.PdfReader(str(path))
    return "\n".join(
        page.extract_text(extraction_mode="layout") or ""
        for page in reader.pages
    )


def _extract_docx_text(path: Path) -> str:
    import docx
    doc = docx.Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


# ---------------------------------------------------------------------------
# Discovery
# ---------------------------------------------------------------------------

async def _fetch_folder_docs(page, folder_id: int, conn, now: str) -> tuple[int, int]:
    """
    Paginate through all documents in a folder via the browser's fetch() API.
    Returns (new_docs_inserted, total_in_folder).
    """
    page_num = 1
    new_docs = 0
    total = 0

    while True:
        try:
            result = await page.evaluate(
                """async (args) => {
                    const resp = await fetch(args.url, {
                        method: 'POST',
                        headers: {
                            'Content-Type': 'application/json',
                            'X-Requested-With': 'XMLHttpRequest',
                        },
                        body: JSON.stringify(args.body),
                    });
                    if (!resp.ok) throw new Error('HTTP ' + resp.status);
                    return resp.json();
                }""",
                {
                    "url": API_PATH,
                    "body": {
                        "folderId": folder_id,
                        "getDocuments": 1,
                        "imageRepo": False,
                        "renderMode": 0,
                        "loadSource": 7,
                        "requestingModuleID": 75,
                        "searchString": "",
                        "pageNumber": page_num,
                        "rowsPerPage": ROWS_PER_PAGE,
                        "sortColumn": "DisplayName",
                        "sortOrder": 0,
                    },
                },
            )
        except Exception as e:
            console.print(f"    [red]API error for folder {folder_id} page {page_num}: {e}[/red]")
            break

        docs = result.get("Documents") or []
        total = result.get("TotalCount", 0)

        for doc in docs:
            doc_id = doc.get("ID")
            if not doc_id:
                continue
            url = doc.get("URL", "")
            if url and not url.startswith("http"):
                url = BASE_URL + url

            cursor = conn.execute(
                """INSERT OR IGNORE INTO documents
                   (id, folder_id, title, url, file_type, published_date, file_size_bytes, discovered_at)
                   VALUES (?,?,?,?,?,?,?,?)""",
                (
                    doc_id,
                    folder_id,
                    doc.get("DisplayName"),
                    url,
                    doc.get("FileType"),
                    doc.get("LastModifiedDateString"),
                    doc.get("FileSize"),
                    now,
                ),
            )
            if cursor.rowcount:
                new_docs += 1

        conn.commit()

        fetched = (page_num - 1) * ROWS_PER_PAGE + len(docs)
        if fetched >= total or not docs:
            break
        page_num += 1

    return new_docs, total


async def _discover_async() -> None:
    init_db()
    conn = get_conn()
    now = _now()

    conn.execute(
        "INSERT OR IGNORE INTO folders (id, parent_id, name, path, discovered_at) VALUES (1,NULL,'DocumentCenter','/',?)",
        (now,),
    )
    conn.commit()

    async with async_playwright() as pw:
        browser = await pw.chromium.launch(headless=True)
        page = await browser.new_page()
        await page.set_extra_http_headers(HEADERS)

        console.print(f"  Loading {DOCENTER_URL} ...")
        await page.goto(DOCENTER_URL, wait_until="networkidle")

        # Expand any collapsible folders in the Ant Design tree
        for _ in range(8):
            switchers = await page.query_selector_all(
                ".ant-tree-switcher:not(.ant-tree-switcher-noop)"
            )
            if not switchers:
                break
            for sw in switchers:
                try:
                    await sw.click()
                    await asyncio.sleep(0.2)
                except Exception:
                    pass
            await page.wait_for_load_state("networkidle")

        # Phase 1: click each folder node to capture its folderId from the POST body
        folder_nodes = await page.query_selector_all(".ant-tree-node-content-wrapper")
        console.print(f"  Found {len(folder_nodes)} folder nodes — clicking to discover IDs...")

        captured: list[int] = []

        def on_request(request):
            if "/Document_AjaxBinding" in request.url and request.method == "POST":
                try:
                    data = json.loads(request.post_data or "{}")
                    fid = data.get("folderId")
                    if fid:
                        captured.clear()
                        captured.append(int(fid))
                except Exception:
                    pass

        page.on("request", on_request)

        folders_found: dict[int, str] = {}  # folder_id -> name
        seen_ids: set[int] = set()

        for node in folder_nodes:
            name = (await node.get_attribute("title") or "").strip()
            if not name:
                continue
            captured.clear()
            try:
                await node.click()
                await asyncio.sleep(0.4)
            except Exception:
                continue
            if captured and captured[0] not in seen_ids:
                fid = captured[0]
                seen_ids.add(fid)
                folders_found[fid] = name
                conn.execute(
                    "INSERT OR IGNORE INTO folders (id, parent_id, name, path, discovered_at) VALUES (?,1,?,?,?)",
                    (fid, name, f"/{name}", now),
                )
                conn.commit()

        page.remove_listener("request", on_request)
        console.print(f"  Discovered {len(folders_found)} folders with IDs")

        # Phase 2: fetch all documents for each folder
        console.print("  Fetching documents...")
        total_new = 0

        for folder_id, folder_name in folders_found.items():
            new_docs, total_in_folder = await _fetch_folder_docs(page, folder_id, conn, now)
            total_new += new_docs
            console.print(
                f"    [{folder_id:>4}] {folder_name:<40} {new_docs:>4} new  ({total_in_folder} total)"
            )

        await browser.close()

    total_docs = conn.execute("SELECT COUNT(*) FROM documents").fetchone()[0]
    total_folders = conn.execute("SELECT COUNT(*) FROM folders").fetchone()[0]
    console.print(
        f"\n  [green]Discovery complete:[/green] "
        f"{total_folders} folders · {total_new} new documents · {total_docs} total in DB"
    )
    conn.close()


def discover() -> None:
    asyncio.run(_discover_async())


# ---------------------------------------------------------------------------
# Download
# ---------------------------------------------------------------------------

def _get_pending_docs(conn, limit, folder_id, retry_errors):
    conditions = ["downloaded_at IS NULL"]
    params: list = []

    if not retry_errors:
        conditions.append("download_error IS NULL")

    if folder_id is not None:
        subtree = _get_folder_subtree(conn, folder_id)
        placeholders = ",".join("?" * len(subtree))
        conditions.append(f"folder_id IN ({placeholders})")
        params.extend(subtree)

    where = " AND ".join(conditions)
    query = f"SELECT id, title, url, file_type FROM documents WHERE {where} ORDER BY id"
    if limit:
        query += f" LIMIT {limit}"

    return conn.execute(query, params).fetchall()


def _get_folder_subtree(conn, folder_id: int) -> list[int]:
    result = [folder_id]
    queue = [folder_id]
    while queue:
        parent = queue.pop()
        for row in conn.execute("SELECT id FROM folders WHERE parent_id=?", (parent,)):
            result.append(row["id"])
            queue.append(row["id"])
    return result


def _extract_text(local_path: Path, file_type: str) -> tuple[str | None, str]:
    if file_type == "pdf":
        try:
            return _extract_pdf_text(local_path), "pypdf"
        except Exception as e:
            console.print(f"    [yellow]PDF extract failed: {e}[/yellow]")
            return None, "none"
    elif file_type in ("docx", "doc"):
        try:
            return _extract_docx_text(local_path), "docx"
        except Exception as e:
            console.print(f"    [yellow]DOCX extract failed: {e}[/yellow]")
            return None, "none"
    return None, "none"


async def _download_one(client: httpx.AsyncClient, doc_id: int, url: str, file_type: str | None) -> dict:
    try:
        resp = await client.get(url, follow_redirects=True, timeout=30)
        resp.raise_for_status()

        ct = resp.headers.get("content-type", "")
        ext = _ext_from_url(str(resp.url)) or _ext_from_content_type(ct)
        if not ext.startswith("."):
            ext = ("." + ext) if ext else ".bin"

        actual_file_type = ext.lstrip(".").lower() or file_type or "bin"
        local_path = DOC_CACHE_DIR / f"{doc_id}{ext}"
        local_path.write_bytes(resp.content)

        text, text_method = _extract_text(local_path, actual_file_type)

        return {
            "local_path": str(local_path),
            "file_type": actual_file_type,
            "file_size_bytes": len(resp.content),
            "error": None,
            "text": text,
            "text_method": text_method,
        }
    except Exception as e:
        return {
            "local_path": None,
            "file_type": None,
            "file_size_bytes": None,
            "error": str(e),
            "text": None,
            "text_method": None,
        }


async def _download_async(limit, folder_id, retry_errors) -> None:
    init_db()
    conn = get_conn()
    now = _now()

    pending = _get_pending_docs(conn, limit, folder_id, retry_errors)
    console.print(f"  {len(pending)} documents to download")
    if not pending:
        conn.close()
        return

    sem = asyncio.Semaphore(DOWNLOAD_CONCURRENCY)

    async def bounded(client, row):
        async with sem:
            result = await _download_one(client, row["id"], row["url"], row["file_type"])
            await asyncio.sleep(DOWNLOAD_DELAY_S)
            return row["id"], result

    ok = err = 0

    async with httpx.AsyncClient(headers=HEADERS) as client:
        with Progress(
            SpinnerColumn(),
            TextColumn("[progress.description]{task.description}"),
            BarColumn(),
            MofNCompleteColumn(),
            console=console,
        ) as progress:
            task = progress.add_task("Downloading...", total=len(pending))
            for coro in asyncio.as_completed([bounded(client, row) for row in pending]):
                doc_id, result = await coro
                progress.advance(task)

                if result["error"]:
                    err += 1
                    conn.execute(
                        "UPDATE documents SET download_error=? WHERE id=?",
                        (result["error"], doc_id),
                    )
                    log_step(conn, doc_id, "download", "error", result["error"])
                else:
                    ok += 1
                    conn.execute(
                        """UPDATE documents
                           SET downloaded_at=?, local_path=?, file_type=?,
                               file_size_bytes=?, download_error=NULL
                           WHERE id=?""",
                        (now, result["local_path"], result["file_type"],
                         result["file_size_bytes"], doc_id),
                    )
                    log_step(conn, doc_id, "download", "ok")

                    if result["text"] is not None:
                        conn.execute(
                            """INSERT OR REPLACE INTO doc_text (doc_id, extracted_at, method, text)
                               VALUES (?,?,?,?)""",
                            (doc_id, now, result["text_method"], result["text"]),
                        )
                        log_step(conn, doc_id, "extract", "ok", result["text_method"])

                conn.commit()

    console.print(f"\n  [green]✓ {ok} downloaded[/green]  [red]{err} errors[/red]")
    conn.close()


def download(limit=None, folder_id=None, retry_errors=False) -> None:
    asyncio.run(_download_async(limit, folder_id, retry_errors))
